"""Parse resume files (PDF, DOCX) into text."""
import os
from pathlib import Path
from typing import Optional
import PyPDF2
import pdfplumber
from docx import Document


class ResumeFileParser:
    """Extract text from resume files."""
    
    @staticmethod
    def parse_pdf_pypdf2(file_path: str) -> str:
        """Parse PDF using PyPDF2 (faster, basic extraction)."""
        text = []
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF with PyPDF2: {e}")
        
        return "\n".join(text)
    
    @staticmethod
    def parse_pdf_pdfplumber(file_path: str) -> str:
        """Parse PDF using pdfplumber (better for complex layouts)."""
        text = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text.append(page_text)
        except Exception as e:
            raise ValueError(f"Failed to parse PDF with pdfplumber: {e}")
        
        return "\n".join(text)
    
    @staticmethod
    def parse_docx(file_path: str) -> str:
        """Parse DOCX file."""
        try:
            doc = Document(file_path)
            text = []
            
            # Extract paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells)
                    if row_text.strip():
                        text.append(row_text)
            
            return "\n".join(text)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {e}")
    
    @staticmethod
    def parse_txt(file_path: str) -> str:
        """Parse plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try different encoding
            with open(file_path, 'r', encoding='latin-1') as file:
                return file.read()
    
    @classmethod
    def parse(cls, file_path: str, method: str = "auto") -> str:
        """
        Parse resume file and extract text.
        
        Args:
            file_path: Path to resume file
            method: Parsing method - "auto", "pypdf2", "pdfplumber"
            
        Returns:
            Extracted text from resume
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_ext = Path(file_path).suffix.lower()
        
        if file_ext == ".pdf":
            if method == "pdfplumber":
                return cls.parse_pdf_pdfplumber(file_path)
            elif method == "pypdf2":
                return cls.parse_pdf_pypdf2(file_path)
            else:  # auto - try pdfplumber first, fallback to pypdf2
                try:
                    return cls.parse_pdf_pdfplumber(file_path)
                except:
                    return cls.parse_pdf_pypdf2(file_path)
        
        elif file_ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path)
        
        elif file_ext == ".txt":
            return cls.parse_txt(file_path)
        
        else:
            raise ValueError(f"Unsupported file format: {file_ext}. Supported: .pdf, .docx, .txt")
    
    @classmethod
    def validate_resume(cls, text: str) -> bool:
        """Basic validation that text looks like a resume."""
        text_lower = text.lower()
        
        # Check for common resume sections
        resume_indicators = [
            'experience', 'education', 'skills', 'work history',
            'employment', 'qualifications', 'resume', 'cv',
            'email', 'phone', 'linkedin'
        ]
        
        found = sum(1 for indicator in resume_indicators if indicator in text_lower)
        
        # Should have at least 2 common resume sections
        return found >= 2 and len(text.split()) >= 50  # At least 50 words
